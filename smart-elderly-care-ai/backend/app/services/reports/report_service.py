"""
report_service.py – Dịch vụ tạo báo cáo y tế (PDF / Excel) từ dữ liệu sinh hiệu và incidents.
"""

import io
from datetime import datetime
from typing import List
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet

from app.models.incident import Incident
from app.models.vital_sign import VitalSign


class ReportService:

    @staticmethod
    def generate_pdf(
        device_id: UUID,
        vitals: List[VitalSign],
        incidents: List[Incident],
    ) -> bytes:
        """Tạo PDF report tóm tắt các chỉ số sinh hiệu và sự kiện cảnh báo."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        elements = []

        # Tiêu đề
        title = Paragraph(f"<b>BÁO CÁO THEO DÕI SỨC KHỎE NGƯỜI CAO TUỔI</b>", styles["Title"])
        subtitle = Paragraph(
            f"Thiết bị Hub ID: {device_id} | Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            styles["Normal"],
        )
        elements.extend([title, subtitle, Spacer(1, 14)])

        # Thống kê tổng quan
        summary_text = Paragraph(
            f"<b>Tổng số bản ghi sinh hiệu:</b> {len(vitals)} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"<b>Số sự kiện cảnh báo:</b> {len(incidents)}",
            styles["Heading3"],
        )
        elements.extend([summary_text, Spacer(1, 10)])

        # Bảng tóm tắt sinh hiệu gần nhất
        if vitals:
            elements.append(Paragraph("<b>Dữ liệu sinh hiệu gần nhất</b>", styles["Heading2"]))
            table_data = [["Thời gian", "Nhịp tim (bpm)", "SpO2 (%)", "Nhiệt độ (°C)", "Té ngã"]]
            for v in vitals[:15]:
                table_data.append([
                    v.time.strftime("%d/%m %H:%M") if v.time else "-",
                    str(v.heart_rate or "-"),
                    str(v.spo2 or "-"),
                    f"{v.skin_temp_max:.1f}" if v.skin_temp_max else "-",
                    "Có" if v.fall_detected else "Không",
                ])
            t = Table(table_data)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2B6CB0")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ]))
            elements.extend([t, Spacer(1, 14)])

        # Bảng sự kiện cảnh báo
        if incidents:
            elements.append(Paragraph("<b>Lịch sử sự kiện cảnh báo</b>", styles["Heading2"]))
            inc_data = [["Thời gian", "Loại cảnh báo", "Mức độ", "Thông điệp", "Đã xử lý"]]
            for inc in incidents[:15]:
                inc_data.append([
                    inc.created_at.strftime("%d/%m %H:%M") if inc.created_at else "-",
                    inc.alert_type,
                    inc.alert_level,
                    inc.message[:30] + ("..." if len(inc.message) > 30 else ""),
                    "Đã duyệt" if inc.is_acknowledged else "Chưa",
                ])
            it = Table(inc_data)
            it.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C53030")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ]))
            elements.append(it)

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_excel(
        device_id: UUID,
        vitals: List[VitalSign],
        incidents: List[Incident],
    ) -> bytes:
        """Tạo bảng tính Excel chứa toàn bộ raw data sinh hiệu và incidents."""
        wb = Workbook()

        # Sheet 1: Vital Signs
        ws_vitals = wb.active
        if ws_vitals is None:
            ws_vitals = wb.create_sheet(title="Vital Signs")
        else:
            ws_vitals.title = "Vital Signs"

        ws_vitals.append(["ID", "Thời gian", "Device ID", "Nhịp tim (bpm)", "SpO2 (%)", "Nhiệt độ (°C)", "Số người", "Té ngã"])
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws_vitals[1]:
            cell.fill = header_fill
            cell.font = header_font

        for v in vitals:
            ws_vitals.append([
                str(v.id),
                v.time.strftime("%Y-%m-%d %H:%M:%S") if v.time else "",
                str(v.device_id),
                v.heart_rate,
                v.spo2,
                v.skin_temp_max,
                v.person_count,
                v.fall_detected,
            ])

        # Sheet 2: Incidents
        ws_inc = wb.create_sheet(title="Incidents")
        ws_inc.append(["ID", "Device ID", "Loại cảnh báo", "Mức độ", "Thông điệp", "Độ tin cậy", "Video URL", "Đã xử lý", "Thời gian"])
        for cell in ws_inc[1]:
            cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
            cell.font = header_font

        for inc in incidents:
            ws_inc.append([
                str(inc.id),
                str(inc.device_id),
                inc.alert_type,
                inc.alert_level,
                inc.message,
                inc.confidence,
                inc.video_clip_url,
                inc.is_acknowledged,
                inc.created_at.strftime("%Y-%m-%d %H:%M:%S") if inc.created_at else "",
            ])

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream.getvalue()


report_service = ReportService()
