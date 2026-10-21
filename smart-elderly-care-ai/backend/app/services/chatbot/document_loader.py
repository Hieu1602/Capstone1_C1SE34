# -*- coding: utf-8 -*-
"""
document_loader.py – Bộ trích xuất và phân đoạn tài liệu đa định dạng cho RAG:
- .pdf: Tài liệu di động (pypdf)
- .docx: Văn bản Microsoft Word (python-docx)
- .xlsx: Bảng tính Microsoft Excel (openpyxl)
- .xls: Bảng tính Microsoft Excel định dạng cũ (xlrd)
- .txt / .md: Văn bản thuần túy và Markdown đa bảng mã

Vị trí: backend/app/services/chatbot/document_loader.py
"""

from __future__ import annotations

import logging
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)


class DocumentLoader:
    """Bộ nạp và xử lý tài liệu đa định dạng cho hệ thống RAG Y tế."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls", ".txt", ".md"}

    def __init__(self, chunk_size: int = 700, chunk_overlap: int = 100):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def load_file(self, file_path: Path | str) -> List[Dict[str, Any]]:
        """Đọc 1 file và trả về danh sách các chunks."""
        path = Path(file_path)
        if not path.is_file() or path.name.startswith("~") or path.name.startswith("."):
            return []

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            logger.debug("[DocLoader] Định dạng không được hỗ trợ: %s", ext)
            return []

        text = ""
        try:
            if ext in {".txt", ".md"}:
                text = self._read_text_file(path)
            elif ext == ".pdf":
                text = self._read_pdf_file(path)
            elif ext == ".docx":
                text = self._read_docx_file(path)
            elif ext == ".xlsx":
                text = self._read_xlsx_file(path)
            elif ext == ".xls":
                text = self._read_xls_file(path)
        except Exception as exc:
            logger.error("[DocLoader] Lỗi đọc file %s (%s): %s", path.name, ext, exc)
            return []

        if not text or len(text.strip()) < 20:
            logger.warning("[DocLoader] File rỗng hoặc không có nội dung chữ: %s", path.name)
            return []

        return self.split_text_into_chunks(text, source_name=path.name)

    def load_directory(self, dir_path: Path | str, recursive: bool = True) -> List[Dict[str, Any]]:
        """Quét và nạp tất cả các file tài liệu trong thư mục."""
        folder = Path(dir_path)
        if not folder.exists() or not folder.is_dir():
            logger.warning("[DocLoader] Thư mục không tồn tại: %s", folder)
            return []

        all_chunks: List[Dict[str, Any]] = []
        pattern = "**/*" if recursive else "*"

        for file_path in folder.glob(pattern):
            if (
                file_path.is_file()
                and not file_path.name.startswith("~")
                and not file_path.name.startswith(".")
                and file_path.suffix.lower() in self.SUPPORTED_EXTENSIONS
            ):
                chunks = self.load_file(file_path)
                all_chunks.extend(chunks)

        logger.info(
            "[DocLoader] Đã nạp thành công %d chunks từ thư mục %s",
            len(all_chunks),
            folder.name,
        )
        return all_chunks

    def split_text_into_chunks(self, text: str, source_name: str) -> List[Dict[str, Any]]:
        """
        Phân đoạn văn bản dựa trên ranh giới tự nhiên (tiêu đề, đoạn văn, câu)
        kèm sliding window overlap để bảo toàn ngữ cảnh y tế liên tục.
        """
        # Chuẩn hoá ký tự xuống dòng và khoảng trắng thừa
        text = re.sub(r"\r\n?", "\n", text)
        text = re.sub(r"[ \t]+", " ", text)

        # Tách theo đoạn văn trước
        raw_paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        if not raw_paragraphs:
            raw_paragraphs = [text.strip()]

        chunks: List[Dict[str, Any]] = []
        current_content: List[str] = []
        current_length = 0

        for para in raw_paragraphs:
            para_len = len(para)

            # Nếu một đoạn quá dài (dài hơn chunk_size), tách nhỏ theo câu
            if para_len > self.chunk_size:
                sub_sentences = re.split(r"(?<=[.!?。])\s+", para)
                for sentence in sub_sentences:
                    s_len = len(sentence)
                    if current_length + s_len > self.chunk_size and current_content:
                        combined = " ".join(current_content).strip()
                        if len(combined) >= 40:
                            chunks.append({
                                "chunk_id": f"{source_name}#chunk_{len(chunks)+1}",
                                "source": source_name,
                                "text": combined,
                            })
                        # Overlap: giữ lại đoạn cuối làm ngữ cảnh nối tiếp
                        overlap_text = combined[-self.chunk_overlap:] if len(combined) > self.chunk_overlap else ""
                        current_content = [overlap_text, sentence] if overlap_text else [sentence]
                        current_length = sum(len(c) for c in current_content)
                    else:
                        current_content.append(sentence)
                        current_length += s_len
            else:
                if current_length + para_len > self.chunk_size and current_content:
                    combined = "\n".join(current_content).strip()
                    if len(combined) >= 40:
                        chunks.append({
                            "chunk_id": f"{source_name}#chunk_{len(chunks)+1}",
                            "source": source_name,
                            "text": combined,
                        })
                    overlap_text = combined[-self.chunk_overlap:] if len(combined) > self.chunk_overlap else ""
                    current_content = [overlap_text, para] if overlap_text else [para]
                    current_length = sum(len(c) for c in current_content)
                else:
                    current_content.append(para)
                    current_length += para_len

        # Đóng gói chunk còn lại
        if current_content:
            combined = "\n".join(current_content).strip()
            if len(combined) >= 30:
                chunks.append({
                    "chunk_id": f"{source_name}#chunk_{len(chunks)+1}",
                    "source": source_name,
                    "text": combined,
                })

        return chunks

    @staticmethod
    def _read_text_file(path: Path) -> str:
        """Đọc file text với fallback bảng mã phổ biến tiếng Việt."""
        encodings = ["utf-8", "utf-8-sig", "cp1258", "latin-1"]
        for enc in encodings:
            try:
                with open(path, "r", encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, LookupError):
                continue
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    @staticmethod
    def _read_pdf_file(path: Path, max_pages: int = 100) -> str:
        """Trích xuất text từ file PDF an toàn từng trang bằng pypdf."""
        from pypdf import PdfReader

        try:
            reader = PdfReader(str(path))
            num_pages = min(len(reader.pages), max_pages)
            extracted: List[str] = []

            for idx in range(num_pages):
                try:
                    page = reader.pages[idx]
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        extracted.append(page_text.strip())
                except Exception as e:
                    logger.warning("[DocLoader] Bỏ qua trang %d lỗi của file %s: %s", idx + 1, path.name, e)

            return "\n\n".join(extracted)
        except Exception as exc:
            logger.error("[DocLoader] Lỗi khi mở file PDF %s: %s", path.name, exc)
            return ""

    @staticmethod
    def _read_docx_file(path: Path) -> str:
        """Trích xuất text từ file Word .docx (gồm đoạn văn và bảng biểu)."""
        import docx

        doc = docx.Document(str(path))
        parts: List[str] = []

        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    parts.append(" | ".join(row_texts))

        return "\n\n".join(parts)

    @staticmethod
    def _read_xlsx_file(path: Path) -> str:
        """Trích xuất dữ liệu từ bảng tính Excel .xlsx (openpyxl)."""
        import openpyxl

        wb = openpyxl.load_workbook(filename=str(path), read_only=True, data_only=True)
        parts: List[str] = []

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            headers: List[str] = []
            sheet_rows: List[str] = []

            for row in sheet.iter_rows(values_only=True):
                if not row or all(v is None or str(v).strip() == "" for v in row):
                    continue

                str_vals = [str(v).strip() if v is not None else "" for v in row]

                if not headers:
                    headers = str_vals
                    sheet_rows.append(f"### Sheet: {sheet_name}")
                    sheet_rows.append("Tiêu đề cột: " + " | ".join(headers))
                else:
                    row_items = []
                    for col_idx, val in enumerate(str_vals):
                        if val:
                            header_name = headers[col_idx] if col_idx < len(headers) and headers[col_idx] else f"Cột {col_idx+1}"
                            row_items.append(f"{header_name}: {val}")
                    if row_items:
                        sheet_rows.append("; ".join(row_items))

            if sheet_rows:
                parts.append("\n".join(sheet_rows))

        wb.close()
        return "\n\n".join(parts)

    @staticmethod
    def _read_xls_file(path: Path) -> str:
        """Trích xuất dữ liệu từ bảng tính Excel định dạng cũ .xls (xlrd)."""
        import xlrd

        wb = xlrd.open_workbook(filename=str(path))
        parts: List[str] = []

        for sheet_idx in range(wb.nsheets):
            sheet = wb.sheet_by_index(sheet_idx)
            if sheet.nrows == 0:
                continue

            headers: List[str] = []
            sheet_rows: List[str] = [f"### Sheet: {sheet.name}"]

            for r in range(sheet.nrows):
                row_vals = sheet.row_values(r)
                str_vals = [str(v).strip() if v is not None else "" for v in row_vals]
                if all(not v for v in str_vals):
                    continue

                if not headers:
                    headers = str_vals
                    sheet_rows.append("Tiêu đề cột: " + " | ".join(headers))
                else:
                    row_items = []
                    for col_idx, val in enumerate(str_vals):
                        if val:
                            header_name = headers[col_idx] if col_idx < len(headers) and headers[col_idx] else f"Cột {col_idx+1}"
                            row_items.append(f"{header_name}: {val}")
                    if row_items:
                        sheet_rows.append("; ".join(row_items))

            if sheet_rows:
                parts.append("\n".join(sheet_rows))

        return "\n\n".join(parts)


# Singleton instance
default_document_loader = DocumentLoader()
